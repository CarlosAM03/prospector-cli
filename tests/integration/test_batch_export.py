"""P95 selected CSV/XLSX exports retain the individual seven-column schema."""

import csv

from openpyxl import load_workbook
import pytest

from models.batch import (
    BatchQuery, BatchQueryResult, BatchSearchResult, DuplicateReference,
    ObservationDisposition, ObservationProvenance, ObservationRef, QueryFailure, QueryStatus,
)
from models.business import Business
from models.normalized_business import NormalizedBusiness
from models.search_query import SearchQuery, Source
from models.search_result import SearchResult
from models.source_identity import VerifiedSourceIdentity
from services.export_service import ExportFormat, ExportService
from engines.batch import select_batch_exports


HEADERS = ["Name", "Category", "Address", "Phone", "Email", "Website", "Language"]
BATCH_ID = "1" * 32


def make_entry(keyword="cafes", export_indices=None):
    query = SearchQuery(Source.GOOGLE_MAPS, keyword, "Tijuana")
    result = SearchResult(
        query,
        businesses=[
            NormalizedBusiness(name="CAFE UNO", phone="664 123 4567", email="a@test.example"),
            NormalizedBusiness(name="CAFE DOS", phone="664 234 5678", email="b@test.example"),
        ],
        original_businesses=[Business("Cafe Uno"), Business("Cafe Dos")],
    )
    return BatchQueryResult(
        BatchQuery(query, 50), QueryStatus.SUCCESS, result=result,
        export_indices=[0, 1] if export_indices is None else export_indices,
        suppressed=[],
    )


@pytest.mark.parametrize("format", [ExportFormat.CSV, ExportFormat.EXCEL])
def test_batch_entry_schema_values_and_originals_are_untouched(tmp_path, monkeypatch, format):
    monkeypatch.chdir(tmp_path)
    entry = make_entry()
    filename = ExportService.export_batch_entry(entry, format, BATCH_ID, 1)
    assert BATCH_ID in filename and "q01" in filename
    path = tmp_path / filename
    if format is ExportFormat.CSV:
        with path.open(newline="", encoding="utf-8") as handle:
            rows = list(csv.reader(handle))
    else:
        workbook = load_workbook(path, read_only=True)
        try:
            assert workbook.sheetnames == ["Businesses"]
            rows = list(workbook["Businesses"].values)
        finally:
            workbook.close()
    assert list(rows[0]) == HEADERS
    assert rows[1][0] == "CAFE UNO"
    assert list(rows[1][3:5]) == ["664 123 4567", "a@test.example"]
    assert len(rows) == 3
    assert entry.result.total_found == 2
    assert entry.result.original_businesses[0].name == "Cafe Uno"


@pytest.mark.parametrize("format", [ExportFormat.CSV, ExportFormat.EXCEL])
def test_fully_suppressed_view_has_headers_and_zero_rows(tmp_path, monkeypatch, format):
    monkeypatch.chdir(tmp_path)
    entry = make_entry()
    entry.export_indices = []
    identity = VerifiedSourceIdentity(Source.GOOGLE_MAPS, "google_place_id", "synthetic")
    entry.suppressed = [
        DuplicateReference(ObservationRef(1, i), ObservationRef(0, i), identity)
        for i in range(2)
    ]
    filename = ExportService.export_batch_entry(entry, format, BATCH_ID, 2)
    path = tmp_path / filename
    if format is ExportFormat.CSV:
        with path.open(newline="", encoding="utf-8") as handle:
            rows = list(csv.reader(handle))
    else:
        workbook = load_workbook(path, read_only=True)
        try:
            rows = list(workbook["Businesses"].values)
        finally:
            workbook.close()
    assert len(rows) == 1 and list(rows[0]) == HEADERS


def test_equal_queries_get_unique_names_and_existing_file_is_not_replaced(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    entry = make_entry()
    first = ExportService.export_batch_entry(entry, ExportFormat.CSV, BATCH_ID, 1)
    second = ExportService.export_batch_entry(entry, ExportFormat.CSV, BATCH_ID, 2)
    assert first != second
    original = (tmp_path / first).read_bytes()
    with pytest.raises(FileExistsError):
        ExportService.export_batch_entry(entry, ExportFormat.CSV, BATCH_ID, 1)
    assert (tmp_path / first).read_bytes() == original


def test_failed_query_never_creates_file(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    request = BatchQuery(SearchQuery(Source.GOOGLE_MAPS, "cafes", "Tijuana"), 50)
    failed = BatchQueryResult(request, QueryStatus.FAILED, error=QueryFailure("navigation", "Safe."))
    with pytest.raises(ValueError, match="failed"):
        ExportService.export_batch_entry(failed, ExportFormat.CSV, BATCH_ID, 1)
    assert list(tmp_path.iterdir()) == []


def test_export_failure_removes_incomplete_file(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)

    def fail_export(self, result, output_path):
        with open(output_path, "w", encoding="utf-8") as handle:
            handle.write("partial")
        raise OSError("disk full")

    monkeypatch.setattr("services.export_service.CsvExporter.export", fail_export)
    with pytest.raises(OSError, match="disk full"):
        ExportService.export_batch_entry(make_entry(), ExportFormat.CSV, BATCH_ID, 1)
    assert list(tmp_path.iterdir()) == []


def test_abc_verified_ids_produce_three_disjoint_csv_files(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    values = [["1", "2", "3"], ["2", "4", "5"], ["1", "5", "6"]]
    entries = []
    for query_index, ids in enumerate(values):
        query = SearchQuery(Source.GOOGLE_MAPS, f"sector {query_index}", "Tijuana")
        result = SearchResult(
            query, [NormalizedBusiness(name=f"ID {value}") for value in ids],
            original_businesses=[Business(f"source {value}") for value in ids],
        )
        entries.append(BatchQueryResult(
            BatchQuery(query, 50), QueryStatus.SUCCESS, result=result,
            export_indices=[0, 1, 2],
            observations=tuple(ObservationProvenance(
                ObservationRef(query_index, position),
                VerifiedSourceIdentity(Source.GOOGLE_MAPS, "google_place_id", value),
                ObservationDisposition.EXPORTED,
            ) for position, value in enumerate(ids)),
        ))
    batch = select_batch_exports(BatchSearchResult(entries))
    paths = [ExportService.export_batch_entry(
        entry, ExportFormat.CSV, BATCH_ID, ordinal,
    ) for ordinal, entry in enumerate(batch.entries, start=1)]
    rows_by_file = []
    for filename in paths:
        with (tmp_path / filename).open(newline="", encoding="utf-8") as handle:
            rows_by_file.append([row[0] for row in list(csv.reader(handle))[1:]])
    assert rows_by_file == [["ID 1", "ID 2", "ID 3"], ["ID 4", "ID 5"], ["ID 6"]]
    assert len(set(paths)) == 3
    assert [entry.result.total_found for entry in batch.entries] == [3, 3, 3]
